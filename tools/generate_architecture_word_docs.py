from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "architecture"
SERVER_DOC = OUTPUT_DIR / "ESP32-Dashboard-Server-Architecture-v1.0.docx"
DEVICE_DOC = OUTPUT_DIR / "ESP32-Dashboard-Device-Runtime-Architecture-v1.0.docx"

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
FONT_MONO = "Consolas"
COLOR_BLUE = "2E74B5"
COLOR_DARK_BLUE = "1F4D78"
COLOR_INK = "1E293B"
COLOR_MUTED = "64748B"
COLOR_LIGHT = "F2F4F7"
COLOR_BLUE_GRAY = "E8EEF5"
COLOR_CALLOUT = "F4F6F9"
COLOR_GREEN = "E8F3EC"
COLOR_AMBER = "FFF4D6"
COLOR_RED = "FCE8E8"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, name: str = FONT_LATIN, east_asia: str = FONT_CJK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "在 Word 中右键更新目录"
    fld_char_3 = OxmlElement("w:fldChar")
    fld_char_3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2, placeholder, fld_char_3])


class ArchitectureDocument:
    def __init__(self, title: str, subtitle: str, document_id: str, review_scope: str):
        self.doc = Document()
        self.title = title
        self.subtitle = subtitle
        self.document_id = document_id
        self.review_scope = review_scope
        self._configure_document()

    def _configure_document(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = FONT_LATIN
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor.from_string(COLOR_INK)
        normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.10

        for name, size, color, before, after in (
            ("Heading 1", 16, COLOR_BLUE, 16, 8),
            ("Heading 2", 13, COLOR_BLUE, 12, 6),
            ("Heading 3", 11.5, COLOR_DARK_BLUE, 8, 4),
        ):
            style = styles[name]
            style.font.name = FONT_LATIN
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(color)
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

        for name in ("List Bullet", "List Number"):
            style = styles[name]
            style.font.name = FONT_LATIN
            style.font.size = Pt(10.5)
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
            style.paragraph_format.left_indent = Inches(0.5)
            style.paragraph_format.first_line_indent = Inches(-0.25)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.10

        if "Code Block" not in styles:
            code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        else:
            code_style = styles["Code Block"]
        code_style.font.name = FONT_MONO
        code_style.font.size = Pt(8.5)
        code_style.font.color.rgb = RGBColor.from_string("263238")
        code_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        code_style.paragraph_format.left_indent = Inches(0.18)
        code_style.paragraph_format.right_indent = Inches(0.18)
        code_style.paragraph_format.space_before = Pt(4)
        code_style.paragraph_format.space_after = Pt(6)
        code_style.paragraph_format.line_spacing = 1.0

        header = section.header
        header_p = header.paragraphs[0]
        header_p.text = f"ESP32-Dashboard | {self.document_id}"
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in header_p.runs:
            set_run_font(run)
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(COLOR_MUTED)

        footer = section.footer
        footer_p = footer.paragraphs[0]
        prefix = footer_p.add_run("Internal Architecture Review  |  ")
        set_run_font(prefix)
        prefix.font.size = Pt(8.5)
        prefix.font.color.rgb = RGBColor.from_string(COLOR_MUTED)
        add_page_field(footer_p)

    def add_cover(self, version: str = "1.0", status: str = "内部评审稿",
                  date: str = "2026-08-20") -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(42)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run("ESP32-DASHBOARD")
        set_run_font(r)
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(COLOR_BLUE)

        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(self.title)
        set_run_font(r)
        r.font.size = Pt(26)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(COLOR_INK)

        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(24)
        r = p.add_run(self.subtitle)
        set_run_font(r)
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor.from_string(COLOR_MUTED)

        table = self.doc.add_table(rows=6, cols=2)
        set_table_geometry(table, [2700, 6660])
        repeat_table_header(table.rows[0])
        rows = [
            ("文档编号", self.document_id),
            ("版本", version),
            ("状态", status),
            ("日期", date),
            ("评审范围", self.review_scope),
            ("设计基线", "配置驱动墨水屏架构 v0.6"),
        ]
        for index, (label, value) in enumerate(rows):
            left, right = table.rows[index].cells
            set_cell_shading(left, COLOR_BLUE_GRAY)
            for cell, text_value, bold in ((left, label, True), (right, value, False)):
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(text_value)
                set_run_font(run)
                run.font.size = Pt(9.5)
                run.font.bold = bold

        self.add_callout(
            "评审目标",
            "确认模块边界、技术栈、数据与安全契约、可靠性目标、迁移路径和第一阶段交付范围。本文是启动实现的架构基线，不替代后续 ADR、接口 Schema 和详细设计。",
            "E8EEF5",
        )
        self.doc.add_page_break()

    def add_toc(self) -> None:
        self.doc.add_heading("目录", level=1)
        p = self.doc.add_paragraph()
        add_toc_field(p)
        self.doc.add_paragraph("提示：首次打开文档后，可在 Word 中更新整个目录和页码。")
        self.doc.add_page_break()

    def heading(self, text: str, level: int = 1) -> None:
        self.doc.add_heading(text, level=level)

    def paragraph(self, text: str, bold_prefix: str | None = None) -> None:
        p = self.doc.add_paragraph()
        if bold_prefix and text.startswith(bold_prefix):
            r1 = p.add_run(bold_prefix)
            set_run_font(r1)
            r1.bold = True
            r2 = p.add_run(text[len(bold_prefix):])
            set_run_font(r2)
        else:
            r = p.add_run(text)
            set_run_font(r)

    def bullets(self, items: Iterable[str], level: int = 0) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            r = p.add_run(item)
            set_run_font(r)

    def numbers(self, items: Iterable[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Number")
            r = p.add_run(item)
            set_run_font(r)

    def code(self, text: str) -> None:
        p = self.doc.add_paragraph(style="Code Block")
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), COLOR_LIGHT)
        p_pr.append(shd)
        for line_index, line in enumerate(text.splitlines()):
            if line_index:
                p.add_run().add_break()
            r = p.add_run(line)
            set_run_font(r, FONT_MONO, FONT_CJK)
            r.font.size = Pt(8.5)

    def add_callout(self, label: str, text: str, fill: str = COLOR_CALLOUT) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        set_table_geometry(table, [TABLE_WIDTH_DXA])
        repeat_table_header(table.rows[0])
        cell = table.cell(0, 0)
        set_cell_shading(cell, fill)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{label}：")
        set_run_font(r)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(COLOR_DARK_BLUE)
        r = p.add_run(text)
        set_run_font(r)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def table(self, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int]) -> None:
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            set_cell_shading(cell, COLOR_BLUE_GRAY)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(header)
            set_run_font(r)
            r.bold = True
            r.font.size = Pt(9)
        repeat_table_header(table.rows[0])
        for row_values in rows:
            row = table.add_row()
            set_repeat_no_split(row)
            for index, value in enumerate(row_values):
                cell = row.cells[index]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(str(value))
                set_run_font(r)
                r.font.size = Pt(8.8)
        set_table_geometry(table, widths)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def figure(self, path: Path, caption: str, width: float = 6.25) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = p.add_run().add_picture(str(path), width=Inches(width))
        shape._inline.docPr.set("title", caption)
        shape._inline.docPr.set("descr", caption)
        caption_p = self.doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_p.paragraph_format.space_after = Pt(8)
        r = caption_p.add_run(caption)
        set_run_font(r)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(COLOR_MUTED)

    def page_break(self) -> None:
        self.doc.add_page_break()

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.core_properties.title = self.title
        self.doc.core_properties.subject = self.subtitle
        self.doc.core_properties.author = "RockBase IoT Architecture Team"
        self.doc.core_properties.keywords = "ESP32, e-paper, dashboard, architecture"
        self.doc.save(path)


def load_font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def rounded_box(draw, box, fill, outline, text, font, radius=18, text_fill="#172033"):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    line_height = font.size + 8 if hasattr(font, "size") else 24
    total_height = len(lines) * line_height
    y = y1 + (y2 - y1 - total_height) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        draw.text(((x1 + x2 - width) / 2, y), line, font=font, fill=text_fill)
        y += line_height


def arrow(draw, start, end, color="#5B6B82", width=5):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - sign * 16, y2 - 10), (x2 - sign * 16, y2 + 10)]
    else:
        sign = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 10, y2 - sign * 16), (x2 + 10, y2 - sign * 16)]
    draw.polygon(points, fill=color)


def create_server_overview(path: Path) -> None:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(34, True)
    box_font = load_font(24, True)
    small_font = load_font(20)
    draw.text((60, 35), "Server Platform Logical Architecture", font=title_font, fill="#1E293B")
    rounded_box(draw, (60, 130, 350, 320), "#E8EEF5", "#2E74B5", "Web Editor\nPlatform Console\nDeveloper Portal", box_font)
    rounded_box(draw, (480, 110, 1120, 350), "#F4F6F9", "#51657A", "API / BFF\nIdentity + Project + Device Profile\nPage + Source + Plugin + Delivery", box_font)
    rounded_box(draw, (1250, 130, 1540, 320), "#E8F3EC", "#2F855A", "ESP32 Devices\nGateway\nLocal Portal", box_font)
    arrow(draw, (350, 225), (480, 225))
    arrow(draw, (1120, 225), (1250, 225))
    rounded_box(draw, (120, 470, 470, 710), "#FFF4D6", "#B7791F", "Connector Runtime\nREST / Webhook\nMQTT / Local Sensor", box_font)
    rounded_box(draw, (625, 470, 975, 710), "#FCE8E8", "#C53030", "Compiler Pipeline\nValidate / Preview\nCompile / Sign", box_font)
    rounded_box(draw, (1130, 470, 1480, 710), "#EDE9FE", "#6B46C1", "Platform Data\nPostgreSQL / Redis\nObject Store / Vault", box_font)
    arrow(draw, (470, 590), (625, 590))
    arrow(draw, (975, 590), (1130, 590))
    arrow(draw, (800, 350), (800, 470))
    draw.text((80, 790), "Modular monolith first; asynchronous workers isolate connectors, rendering and package compilation.", font=small_font, fill="#64748B")
    image.save(path)


def create_package_pipeline(path: Path) -> None:
    image = Image.new("RGB", (1700, 580), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(32, True)
    box_font = load_font(21, True)
    draw.text((50, 30), "Immutable Display Package Publication Pipeline", font=title_font, fill="#1E293B")
    labels = [
        "Project\nDraft",
        "Schema +\nBudget Validate",
        "Pixel Accurate\nPreview",
        "Compile\nResources",
        "Sign +\nStore",
        "Device\nDelivery",
        "Activate +\nAck",
    ]
    fills = ["#E8EEF5", "#F4F6F9", "#EDE9FE", "#FFF4D6", "#FCE8E8", "#E8F3EC", "#E8EEF5"]
    x = 45
    boxes = []
    for label, fill in zip(labels, fills):
        box = (x, 190, x + 190, 360)
        rounded_box(draw, box, fill, "#51657A", label, box_font, radius=16)
        boxes.append(box)
        x += 235
    for left, right in zip(boxes, boxes[1:]):
        arrow(draw, (left[2], 275), (right[0], 275), width=4)
    draw.text((65, 435), "candidate never overwrites active before hash, signature, profile, driver, font and resource checks pass", font=load_font(20), fill="#64748B")
    image.save(path)


def create_device_layers(path: Path) -> None:
    image = Image.new("RGB", (1500, 1000), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(34, True)
    box_font = load_font(24, True)
    small_font = load_font(19)
    draw.text((55, 30), "EInk Dashboard Runtime Layers", font=title_font, fill="#1E293B")
    layers = [
        ("Package Runtime", "DeviceProfileVerifier / PackageStore / PageManager / Scheduler", "#E8EEF5"),
        ("Application Runtime", "WidgetRegistry / BindingResolver / FocusSessionService / AppState", "#F4F6F9"),
        ("Data Runtime", "SnapshotStore / Cache / Freshness / LocalSensorBridge", "#FFF4D6"),
        ("Graphics Runtime", "LayoutEngine / FontEngine / ColorMapper / Compositor / RefreshPlanner", "#EDE9FE"),
        ("System Services", "Network / Time / RTC / Input / Power / Storage / OTA / Diagnostics", "#E8F3EC"),
        ("Hardware Abstraction", "BSP / DisplayDriverAdapter / SensorAdapter / ButtonAdapter", "#FCE8E8"),
    ]
    y = 120
    for name, detail, fill in layers:
        draw.rounded_rectangle((80, y, 1420, y + 115), radius=16, fill=fill, outline="#51657A", width=3)
        draw.text((115, y + 22), name, font=box_font, fill="#1E293B")
        draw.text((430, y + 27), detail, font=small_font, fill="#334155")
        y += 135
    draw.text((300, 935), "Arduino-ESP32 / ESP-IDF + FreeRTOS + GxEPD2 + Hardware Drivers", font=box_font, fill="#64748B")
    image.save(path)


def create_device_state_machine(path: Path) -> None:
    image = Image.new("RGB", (1700, 720), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(32, True)
    box_font = load_font(20, True)
    draw.text((50, 25), "Wake-Render-Sleep Control Flow", font=title_font, fill="#1E293B")
    labels = [
        "Boot / Wake",
        "Verify Hardware\nProfile",
        "Load Active /\nCandidate Package",
        "Resolve Page +\nDependencies",
        "Acquire\nSnapshots",
        "Layout +\nCompose",
        "Refresh +\nCommit",
        "Interactive or\nDeep Sleep",
    ]
    positions = [(40, 150), (250, 150), (500, 150), (775, 150), (1050, 150), (1290, 150), (1290, 430), (970, 430)]
    boxes = []
    for (x, y), label in zip(positions, labels):
        width = 200 if x < 500 else 225
        box = (x, y, x + width, y + 145)
        rounded_box(draw, box, "#F4F6F9", "#2E74B5", label, box_font, radius=14)
        boxes.append(box)
    for left, right in zip(boxes[:6], boxes[1:6]):
        arrow(draw, (left[2], 222), (right[0], 222), width=4)
    arrow(draw, (boxes[5][2], 222), (boxes[6][0], 500), width=4)
    arrow(draw, (boxes[6][0], 500), (boxes[7][2], 500), width=4)
    arrow(draw, (boxes[7][0], 560), (570, 650), width=4)
    draw.text((115, 625), "Button / Schedule / Package / Snapshot / Focus / Power events return through one queue", font=load_font(19), fill="#64748B")
    image.save(path)


def add_common_document_control(doc: ArchitectureDocument) -> None:
    doc.heading("1. 文档控制与评审方法", 1)
    doc.paragraph("本文用于技术评审和项目启动。涉及协议、数据结构和接口的示例是建议基线，实施前应固化为 JSON Schema、OpenAPI、数据库迁移和 ADR。")
    doc.table(
        ["评审角色", "主要关注点", "预期输出"],
        [
            ["产品/设计", "设备选择、页面编排、数据源配置和错误状态是否完整", "范围确认、交互缺口"],
            ["平台架构", "领域边界、数据一致性、扩展性和部署复杂度", "架构批准或 ADR"],
            ["固件架构", "Profile、包格式、资源预算、运行状态机和兼容性", "设备契约批准"],
            ["安全/隐私", "身份、密钥、插件权限、日志脱敏和数据保留", "威胁模型与整改项"],
            ["测试/交付", "可测性、发布门禁、回滚和真机验证", "验收矩阵与里程碑"],
        ],
        [1500, 4300, 3560],
    )
    doc.add_callout("评审原则", "先冻结跨端契约，再并行实现平台和固件。任何会影响页面包兼容性的变更都必须通过版本化 Schema 和迁移策略处理。")


def build_server_document(asset_dir: Path) -> None:
    overview = asset_dir / "server_overview.png"
    pipeline = asset_dir / "package_pipeline.png"
    create_server_overview(overview)
    create_package_pipeline(pipeline)

    d = ArchitectureDocument(
        "服务器端架构设计",
        "配置驱动墨水屏平台、编辑器、数据连接器与页面包交付体系",
        "RB-EPD-ARCH-SERVER-001",
        "平台服务、Web 编辑器、数据源、插件生态、设备管理、发布和运维",
    )
    d.add_cover()
    d.add_toc()
    add_common_document_control(d)

    d.heading("2. 目标、范围与关键决策", 1)
    d.heading("2.1 建设目标", 2)
    d.bullets([
        "用户在创建 Dashboard 项目时先选择 Device Profile，编辑器随后锁定分辨率、物理尺寸、色板、驱动、刷新、字体和资源预算。",
        "页面、控件、数据绑定、排序和主页均由服务器管理，发布为不可变且可回滚的多页 Display Package。",
        "数据源通过自描述 SourceDefinition、SourceInstance、Connector Runtime 和 Normalized Snapshot 解耦。",
        "平台能够支持公共云和自托管部署，并允许设备在平台离线时继续使用最后有效页面包和缓存。",
        "第一阶段以 NM-EPD-420 400 x 300 黑白与黑白红设备为验收目标，架构支持后续多分辨率和多色屏。",
    ])
    d.heading("2.2 非目标", 2)
    d.bullets([
        "不向 ESP32 下发 JavaScript、Lua、动态库或未经审核的原生代码。",
        "不在设备本地网页实现页面画布、排序或主页编辑。",
        "不以 TRMNL 协议兼容或接入其设备为目标。",
        "第一阶段不建设复杂计费系统、完整插件商业市场和多区域容灾。",
        "不承诺把一个页面草稿无损迁移到任意设备；设备迁移必须生成新的 Profile 变体并重新校验。",
    ])
    d.heading("2.3 已冻结的架构决策", 2)
    d.table(
        ["决策", "结论", "原因"],
        [
            ["服务形态", "模块化单体 + 异步 Worker", "降低初期运维成本，同时保留清晰领域边界"],
            ["主要语言", "TypeScript", "共享 API 类型、Schema、Web 编辑器和服务端工具链"],
            ["Web", "Next.js + React + TypeScript", "成熟的路由、SSR、组件生态和部署方式"],
            ["API", "Fastify + TypeBox/OpenAPI", "轻量、性能可控、Schema 驱动，适合后续拆分"],
            ["数据", "PostgreSQL + Redis + S3/MinIO", "事务数据、队列/缓存、不可变制品分别治理"],
            ["预览核心", "可移植 C++17 核心编译为 WASM/native", "与 ESP32 共享布局、字体 metrics 和颜色量化规则"],
            ["设备交付", "HTTPS Pull + 短期令牌", "适合低功耗设备唤醒拉取和断点重试"],
            ["事件", "Transactional Outbox + BullMQ", "避免首期引入 Kafka，同时保证可靠异步处理"],
        ],
        [1500, 2900, 4960],
    )

    d.heading("3. 系统上下文与逻辑架构", 1)
    d.paragraph("平台服务四类主体：终端用户、插件/设备开发者、运维管理员和 ESP32 设备。设备通过 HTTPS 完成配对、状态上报、页面包检查和下载；浏览器不直接向设备下发页面定义。")
    d.figure(overview, "图 1 服务器端逻辑架构")
    d.heading("3.1 逻辑边界", 2)
    d.table(
        ["边界", "负责内容", "明确不负责"],
        [
            ["Platform Console", "账户、设备、项目、页面、数据源、插件、发布、诊断", "直接访问设备私钥或渲染硬件"],
            ["Editor Core", "画布、属性、绑定、Profile 预览、预算和校验", "持久化密钥、执行任意插件代码"],
            ["Platform API", "领域事务、授权、Schema、查询和命令", "长耗时抓取、图片渲染和包编译"],
            ["Connector Runtime", "拉取/接收/标准化数据，形成 Snapshot", "页面布局和设备刷新"],
            ["Compiler Runtime", "校验、预览、资源编译、签名和制品写入", "修改已发布版本"],
            ["Delivery Runtime", "设备检查更新、签发下载 URL、记录确认和回滚", "替设备决定局刷/全刷"],
        ],
        [1800, 3700, 3860],
    )

    d.heading("4. 技术栈与代码组织", 1)
    d.heading("4.1 推荐技术基线", 2)
    d.table(
        ["层", "技术", "实施说明"],
        [
            ["Monorepo", "pnpm workspace + Turborepo", "共享 Schema、SDK、UI 和构建缓存；版本统一发布"],
            ["Web", "Next.js 15 / React / TypeScript", "App Router；编辑器区域主要使用 Client Components"],
            ["状态管理", "TanStack Query + Zustand", "服务端状态与画布瞬态状态分离"],
            ["画布", "原生 Canvas/WebGL + 自研场景图", "不使用 HTML/CSS 作为最终渲染真值"],
            ["API", "Fastify + TypeBox + OpenAPI", "请求、响应和数据库 DTO 从 Schema 派生"],
            ["数据库", "PostgreSQL 16 + Drizzle/Prisma", "建议优先 Drizzle，迁移 SQL 可审计"],
            ["队列", "Redis 7 + BullMQ", "Connector、Preview、Compile、Delivery、Notification 队列"],
            ["对象存储", "S3 兼容存储", "页面包、字体包、插件资源、预览图和诊断包"],
            ["密钥", "Cloud KMS/Vault；开发用 SOPS", "数据库只保存 key_ref 和加密元数据"],
            ["可观测性", "OpenTelemetry + Prometheus + Loki", "traceId/jobId/packageId/deviceId 贯穿"],
        ],
        [1550, 2700, 5110],
    )
    d.heading("4.2 推荐目录", 2)
    d.code("""apps/
  web/                 # Platform Console + Page Editor
  api/                 # Fastify modular monolith
  worker/              # Connector / Preview / Compile jobs
packages/
  contracts/           # TypeBox JSON Schema + generated types
  editor-core/         # scene graph, commands, undo/redo
  renderer-wasm/       # portable C++ renderer WASM binding
  device-profiles/     # versioned profiles and fixtures
  plugin-sdk/          # manifest, recipe and source validation
  ui/                  # RockBaseIOT VI components
services/
  renderer-native/     # deterministic reference rendering CLI
  signing/             # isolated signing adapter
infra/
  compose/ helm/ terraform/ migrations/""")
    d.heading("4.3 模块化单体拆分原则", 2)
    d.bullets([
        "每个领域模块拥有自己的应用服务、仓储接口、事件和 API 路由，不允许跨模块直接访问数据表。",
        "共享 contracts 只包含稳定 DTO、事件和 Schema，不放业务实现。",
        "Connector、Preview、Compile 和签名属于资源密集或安全敏感任务，通过队列隔离。",
        "达到独立扩缩容、独立故障域或独立合规要求后，再把模块拆成服务。",
    ])

    d.heading("5. 领域模型与数据所有权", 1)
    d.table(
        ["聚合", "核心对象", "所有权与不变量"],
        [
            ["Identity", "Account, Workspace, Membership", "用户和租户边界；所有资源必须归属 Workspace"],
            ["Device", "Device, Pairing, DeviceCredential", "Device UUID 唯一；密钥不可通过普通 API 返回"],
            ["Profile", "DeviceProfile, DriverBinding, FontPack", "版本不可变；发布后只新增版本"],
            ["Project", "DashboardProject, ProjectVariant", "草稿绑定唯一 targetProfileId"],
            ["Page", "PageDraft, PageNode, PageOrder", "主页唯一且 pageNo=0；已发布页面不可变"],
            ["Source", "SourceDefinition, SourceInstance, SecretRef", "定义可共享，实例和凭据归属租户"],
            ["Snapshot", "NormalizedSnapshot, Freshness", "按 SourceInstance 和版本保存；可设置 TTL"],
            ["Plugin", "PluginVersion, RecipeVersion, Permission", "版本不可变，权限显式声明"],
            ["Release", "DisplayPackage, Artifact, Delivery", "candidate/active/previous 可追溯，hash 唯一"],
        ],
        [1550, 3000, 4810],
    )
    d.heading("5.1 关键关系", 2)
    d.code("""Workspace 1---N Device
Workspace 1---N DashboardProject
DashboardProject 1---N ProjectVariant
ProjectVariant 1---N PageDraft
ProjectVariant N---1 DeviceProfile
PageDraft N---N SourceInstance (through Binding)
SourceInstance 1---N NormalizedSnapshot
ProjectVariant 1---N DisplayPackage
DisplayPackage 1---N DeviceDelivery""")
    d.heading("5.2 页面顺序和主页规则", 2)
    d.bullets([
        "pageId 是稳定业务 ID；pageNo 是发布时生成的连续显示索引。",
        "任意页面可设为主页，发布器把该页调整为 pageNo=0。",
        "删除、禁用或拖动页面只影响草稿；已发布包保持不变。",
        "设备只可修改当前游标、轮转开关和间隔，不可写回页面布局与顺序。",
    ])

    d.heading("6. Device Profile、字体与精确预览", 1)
    d.heading("6.1 设计入口", 2)
    d.numbers([
        "创建项目并选择设备型号及 Device Profile。",
        "平台锁定像素尺寸、有效物理尺寸、DPI、色板、Driver、刷新策略、Font Pack、模板和内存预算。",
        "编辑器只展示该 Profile 支持的控件、字体角色、颜色和模板。",
        "发布时写入 profileId、firmwareTarget、driverId、Driver Adapter 版本和字体 hashes。",
        "设备激活前执行同一组兼容性检查。",
    ])
    d.heading("6.2 Profile Registry", 2)
    d.table(
        ["字段组", "必要字段", "平台用途"],
        [
            ["Identity", "profileId, deviceModel, boardRevision, firmwareTarget", "项目绑定、固件筛选和迁移"],
            ["Display", "width/height px, active mm, orientation, safeArea", "原生画布、DPI 与物理预览"],
            ["Driver", "driverId, family, controller, adapterVersion", "包兼容和固件 Target 约束"],
            ["Color", "mode, palette, logical mapping, dithering", "控件筛选和色板量化"],
            ["Refresh", "full/partial/region, interval, ghosting budget", "发布警告和刷新策略模拟"],
            ["Memory", "heap, PSRAM, package, node, chart limits", "实时预算和发布门禁"],
            ["Font", "allowed packs, glyph budget, metrics hashes", "精确测量和资源编译"],
            ["Capability", "buttons, sensors, RTC, network, power", "控件和交互能力过滤"],
        ],
        [1500, 4550, 3310],
    )
    d.heading("6.3 渲染一致性", 2)
    d.paragraph("推荐把布局、文本测量、字形 atlas、逻辑颜色映射和抖动实现为不依赖 Arduino 的 C++17 核心。该核心分别编译为 ESP32 静态库、服务器 native CLI 和浏览器 WASM。GxEPD2 只位于设备最终输出适配层。")
    d.bullets([
        "Pixel Accurate 画布始终使用设备原生像素，不使用 CSS 自动缩放作为真值。",
        "字体从 Inter、Noto Sans CJK、Source Han Sans、Roboto Mono 等免费字体生成统一 glyph 资产。",
        "Physical Size 模式根据 active mm 和浏览器标尺校准，目标误差为 X/Y 各不超过 +/-2%。",
        "CI 对相同 Profile、Font Pack、Snapshot 和页面 fixture 比较色板量化后的像素 hash。",
    ])

    d.heading("7. Web 图形化编辑器", 1)
    d.heading("7.1 编辑器组成", 2)
    d.table(
        ["区域", "功能", "关键实现"],
        [
            ["Project Shell", "设备、草稿、发布版本和协作上下文", "路由级加载与权限校验"],
            ["Page Navigator", "创建、复制、删除、禁用、拖动、设置主页", "命令模式 + 乐观更新 + 服务端版本号"],
            ["Canvas", "选择、拖动、缩放、对齐、网格、图层", "场景图；坐标固定为设备原生像素"],
            ["Inspector", "控件属性、字体角色、颜色角色、刷新策略", "由 Widget Schema 自动生成表单"],
            ["Binding Editor", "选择 SourceInstance 和字段路径", "由 Source Schema 过滤兼容控件"],
            ["Preview", "Fresh/Stale/Empty/Error、语言、Profile", "WASM renderer + fixture"],
            ["Budget Panel", "Flash、Heap、PSRAM、节点、图表点", "增量估算；超限阻止发布"],
        ],
        [1500, 3300, 4560],
    )
    d.heading("7.2 编辑命令与并发", 2)
    d.bullets([
        "前端使用 command pattern 实现 undo/redo，命令只改变规范化场景树。",
        "草稿保存采用 revision 或 ETag 的乐观并发控制，冲突时返回 409 和字段级差异。",
        "自动保存只写草稿；发布需要显式用户操作和完整校验。",
        "第一阶段不做多人实时协同；保留 presence 和 CRDT 接口作为后续能力。",
    ])

    d.heading("8. 数据源与 Connector Runtime", 1)
    d.heading("8.1 三层模型", 2)
    d.code("""SourceDefinition (plugin/platform owned)
  -> SourceInstance (workspace owned, binds key_ref and parameters)
  -> NormalizedSnapshot (versioned data contract consumed by widgets)""")
    d.heading("8.2 协议执行器", 2)
    d.table(
        ["协议", "运行位置", "安全与可靠性要求"],
        [
            ["rest_poll", "平台 Worker 或设备", "HTTPS、SSRF 防护、超时、重试、限流、ETag"],
            ["webhook_push", "平台入口", "签名、幂等键、时间窗、防重放、速率限制"],
            ["mqtt", "平台或 Gateway", "独立凭据、Topic ACL、QoS、保留消息策略"],
            ["local_sensor", "设备采样并上报", "设备身份、采样时间、单位、离线与批量上报"],
        ],
        [1500, 2400, 5460],
    )
    d.heading("8.3 调度和缓存", 2)
    d.bullets([
        "调度键为 workspace/sourceInstance，使用分布式锁避免重复抓取。",
        "连接器先写原始响应摘要和解析结果，再原子发布新 Snapshot；失败不覆盖最后有效数据。",
        "每个 SourceDefinition 声明 interval、cache_ttl、stale_after、rate_limit 和 retry policy。",
        "页面发布器静态计算依赖清单；设备只请求当前页所需 Snapshot。",
        "敏感原始载荷默认不长期保存；用于调试时需要用户授权、脱敏和短 TTL。",
    ])
    d.heading("8.4 Freshness 契约", 2)
    d.code("""{
  "snapshotVersion": 12,
  "observedAt": "2026-08-20T09:00:00Z",
  "fetchedAt": "2026-08-20T09:00:03Z",
  "generatedAt": "2026-08-20T09:00:04Z",
  "expiresAt": "2026-08-20T10:00:00Z",
  "sourceState": "ready | stale | expired | error",
  "data": { }
}""")

    d.heading("9. 插件、Recipe 与 SDK", 1)
    d.table(
        ["对象", "内容", "执行边界"],
        [
            ["Widget", "固件内置控件类型和稳定 Schema", "设备执行受限渲染逻辑"],
            ["Recipe", "页面布局、控件、绑定、状态和 Profile 变体", "平台编译为页面包"],
            ["Plugin", "Manifest、SourceDefinition、Recipe、资源、文档和 fixtures", "平台校验、签名和分发；不下发原生代码"],
        ],
        [1450, 4050, 3860],
    )
    d.heading("9.1 插件供应链", 2)
    d.numbers([
        "上传插件版本并计算内容 hash。",
        "校验 Manifest、许可证、资源路径、大小、权限和 Profile 兼容。",
        "在隔离 Worker 中运行静态检查和 fixture 渲染。",
        "生成 SBOM、兼容性报告和安全审查状态。",
        "批准后发布不可变版本；撤销只影响新安装，已安装版本给出升级/回滚提示。",
    ])
    d.add_callout("第一阶段边界", "只建设标准插件仓库和管理员发布流程，不开放任意社区插件自动上架。", COLOR_AMBER)

    d.heading("10. 页面编译、签名与发布", 1)
    d.figure(pipeline, "图 2 Display Package 发布流水线")
    d.heading("10.1 编译阶段", 2)
    d.numbers([
        "加载冻结的 ProjectVariant、Device Profile、Widget/Source Schema 和 Font Pack。",
        "执行结构、权限、字段绑定、布局边界、字体、颜色、刷新和内存预算校验。",
        "用 fixture 或最新 Snapshot 生成所有状态预览并计算像素 hash。",
        "裁剪字体字形、图标和静态资源，生成确定性 payload。",
        "生成 Display Envelope、payload hash、SBOM 和兼容性清单。",
        "调用隔离 Signing Service 签名，并写入对象存储。",
        "事务提交 DisplayPackage 和 outbox 事件，交由 Delivery Worker 分发。",
    ])
    d.heading("10.2 包格式关键字段", 2)
    d.code("""{
  "packageId": "dashboard-20260820-001",
  "packageVersion": 12,
  "deviceProfile": "nm-epd-420-bwr-400x300-v1",
  "firmwareTarget": "nm-display-420-bwr",
  "driverId": "gxepd2.gdey042z98.bwr",
  "driverAdapterVersion": 1,
  "profileSchemaVersion": 1,
  "runtimeSchemaVersion": 1,
  "widgetApiVersion": 1,
  "fontPackId": "font.noto-inter.400x300.v1",
  "fontPackSha256": "...",
  "fontMetricsSha256": "...",
  "sha256": "...",
  "signature": "..."
}""")
    d.heading("10.3 原子性和回滚", 2)
    d.bullets([
        "草稿、编译任务和已发布包是不同对象；失败任务不得污染已发布版本。",
        "包版本不可变；回滚通过重新指向历史版本，不修改原包。",
        "制品上传完成且 hash 校验通过后，才提交数据库可见状态。",
        "签名密钥由 KMS/HSM 管理，API 和 Worker 不读取明文私钥。",
    ])

    d.heading("11. 设备管理、配对与交付", 1)
    d.heading("11.1 身份和配对", 2)
    d.code("""Device first boot
  -> generate Device UUID + P-256 key pair
  -> request short-lived claim challenge
  -> display pairing code / QR
  -> user confirms device model and code
  -> device signs challenge
  -> platform binds public key and issues rotating credential""")
    d.heading("11.2 设备 API", 2)
    d.table(
        ["接口", "用途", "约束"],
        [
            ["POST /v1/device/claim", "申请一次性配对挑战", "限流；不信任 MAC"],
            ["POST /v1/device/session", "签名挑战并换取短期令牌", "P-256；令牌可轮换"],
            ["GET /v1/device/package", "检查当前目标包", "ETag；可返回 no-change"],
            ["GET /v1/artifacts/{id}", "断点下载页面包/OTA", "短期签名 URL；Range"],
            ["POST /v1/device/ack", "上报下载、激活或回滚结果", "不包含用户页面内容"],
            ["POST /v1/device/telemetry", "可选脱敏运行指标", "用户可关闭；批量"],
        ],
        [2450, 3150, 3760],
    )
    d.heading("11.3 低功耗交付策略", 2)
    d.bullets([
        "设备唤醒时先发送 firmware/profile/packageVersion 和所需 Snapshot 版本摘要。",
        "服务端返回 no-change、package、snapshot-delta、ota 或 recovery 指令。",
        "下载支持 HTTP Range、ETag 和校验；设备确认激活后更新 Delivery 状态。",
        "平台不能强制设备立即刷新屏幕，最终 RefreshPlanner 由设备决定。",
    ])

    d.heading("12. API、事件和版本治理", 1)
    d.heading("12.1 API 原则", 2)
    d.bullets([
        "外部 REST API 使用 /v1，OpenAPI 是契约真值。",
        "命令请求支持 Idempotency-Key；更新资源使用 revision/If-Match。",
        "错误返回稳定 code、message、traceId 和可选 fieldErrors。",
        "批量接口显式限制条数和响应大小；大文件只通过对象存储 URL 传输。",
    ])
    d.heading("12.2 关键领域事件", 2)
    d.table(
        ["事件", "生产者", "消费者"],
        [
            ["ProjectPublished", "Project", "Compiler"],
            ["SnapshotUpdated", "Connector", "Preview cache、Delivery"],
            ["PackageBuilt", "Compiler", "Signing、Delivery"],
            ["PackageActivated", "Device API", "Device registry、analytics"],
            ["DeviceProfileReleased", "Profile registry", "Compatibility index、Editor"],
            ["PluginVersionApproved", "Plugin registry", "Catalog、Compiler cache"],
        ],
        [2600, 2600, 4160],
    )
    d.heading("12.3 Schema 兼容规则", 2)
    d.bullets([
        "增加可选字段属于 minor 兼容变更；删除字段或改变语义提升 major。",
        "所有已发布包保留其 Schema 和资源版本；编译器可读取当前和受支持的历史版本。",
        "设备能力、Widget API、Font Pack 和包 Schema 分别版本化，不使用单一全局版本掩盖兼容性。",
    ])

    d.heading("13. 存储设计", 1)
    d.heading("13.1 PostgreSQL", 2)
    d.table(
        ["表组", "代表表", "关键索引/约束"],
        [
            ["Identity", "workspaces, users, memberships", "workspace_id + subject 唯一"],
            ["Device", "devices, device_keys, pairings, deliveries", "device_uuid 唯一；公钥指纹唯一"],
            ["Profile", "device_profiles, font_packs, compatibility", "profile_id + version 唯一"],
            ["Project", "projects, variants, page_drafts, bindings", "variant revision；主页唯一部分索引"],
            ["Source", "source_definitions, source_instances, snapshots", "instance + snapshot_version；TTL 索引"],
            ["Plugin", "plugins, plugin_versions, reviews", "plugin_id + semver 唯一"],
            ["Release", "packages, artifacts, package_deliveries", "package hash 唯一；状态索引"],
            ["Audit", "audit_events, outbox_events", "workspace + created_at；outbox 状态"],
        ],
        [1500, 3700, 4160],
    )
    d.heading("13.2 Redis 与对象存储", 2)
    d.bullets([
        "Redis 用于队列、短期锁、速率限制、会话和可丢弃缓存，不保存唯一业务真值。",
        "对象存储保存 immutable package、preview、font/icon/plugin bundle、SBOM 和诊断包。",
        "对象 key 包含 workspace 隔离前缀和内容 hash；私有对象默认禁止公开列表。",
        "生命周期策略分别管理原始响应、预览、历史包和诊断包。",
    ])

    d.heading("14. 安全与隐私", 1)
    d.table(
        ["威胁", "控制措施", "验证"],
        [
            ["设备冒充", "设备 UUID + P-256 私钥、签名挑战、短期令牌", "配对和重放测试"],
            ["SSRF", "域名/IP 策略、DNS 重绑定防护、禁止私网段、出口代理", "恶意 URL 用例"],
            ["Webhook 重放", "签名、时间窗、nonce/事件 ID 幂等", "重复事件测试"],
            ["密钥泄露", "Vault/KMS、key_ref、字段加密、日志脱敏", "Secret scanning"],
            ["插件供应链", "不可变版本、hash、签名、资源白名单、SBOM", "隔离构建和审计"],
            ["租户越权", "所有查询携带 workspace scope；对象存储短期 URL", "权限矩阵测试"],
            ["页面包篡改", "hash、签名、profile/driver/font 绑定", "设备负向激活测试"],
        ],
        [1850, 4700, 2810],
    )
    d.heading("14.1 隐私默认值", 2)
    d.bullets([
        "默认不采集 WiFi 密码、Secret ICS URL、API Key、日历正文、持仓明细和精确位置。",
        "诊断包由用户主动导出，先脱敏并显示将包含的字段。",
        "页面内容和 Snapshot 不用于产品分析；质量分析优先使用版本、计数和错误类别。",
        "数据保留策略按 SourceDefinition 和 Workspace 配置，支持删除与导出。",
    ])

    d.heading("15. 可靠性、可观测性与容量", 1)
    d.heading("15.1 服务目标", 2)
    d.table(
        ["能力", "第一阶段目标", "降级行为"],
        [
            ["项目读写 API", "月可用性 99.9%", "只读查看最近草稿"],
            ["设备包检查", "月可用性 99.9%", "设备继续 active 包"],
            ["Connector", "按源统计成功率和延迟", "继续 last-good Snapshot + stale"],
            ["编译任务", "P95 小于 30 秒", "重试或回退上一已发布包"],
            ["对象下载", "支持 Range；hash 100% 校验", "断点重试，不激活 candidate"],
        ],
        [2200, 3000, 4160],
    )
    d.heading("15.2 可观测性", 2)
    d.bullets([
        "所有请求和任务携带 traceId；编译链增加 projectId/packageId，设备链增加 deviceId。",
        "指标覆盖 API 延迟、队列深度、Connector 失败分类、缓存命中、编译耗时、包大小和设备激活率。",
        "日志使用结构化 JSON，敏感字段在 Logger 层统一脱敏。",
        "告警以用户影响为中心：发布失败率、设备激活下降、Source 大面积失效和签名服务不可用。",
    ])
    d.heading("15.3 容量假设", 2)
    d.paragraph("启动容量建议按 10,000 台设备、每台每日 48 次唤醒、平均 5 个启用页面、每个页面 2 个 Snapshot 依赖估算。设备包检查与数据抓取必须解耦，不能按每次设备唤醒重复抓取第三方数据。")

    d.heading("16. 部署与运维拓扑", 1)
    d.heading("16.1 环境", 2)
    d.table(
        ["环境", "拓扑", "用途"],
        [
            ["Local", "Docker Compose + MinIO + PostgreSQL + Redis", "个人开发、插件和 Profile fixture"],
            ["Integration", "共享集成环境 + 模拟设备群", "跨端契约和升级测试"],
            ["Staging", "与生产同构，小规模真实设备", "签名、OTA、回滚、性能和安全验证"],
            ["Production", "托管数据库/Redis/对象存储 + 容器平台", "水平扩展 API/Worker，签名隔离"],
        ],
        [1500, 3900, 3960],
    )
    d.heading("16.2 发布策略", 2)
    d.bullets([
        "数据库迁移采用 expand/contract，两阶段兼容。",
        "API、Worker 和 Compiler 使用同一 contracts 版本矩阵。",
        "设备 Package 和 OTA 分批灰度，按 Device Profile、固件版本和测试组控制。",
        "每次发布保留应用版本、Schema、Compiler、Font Pack 和签名 key version 关联。",
    ])

    d.heading("17. 测试与 CI/CD", 1)
    d.table(
        ["测试层", "内容", "门禁"],
        [
            ["Schema", "Profile/Plugin/Recipe/Source/Package 正负 fixtures", "兼容性和迁移全部通过"],
            ["Domain", "主页、排序、引用、权限、状态机", "单元和属性测试"],
            ["Connector", "HTTP、RSS、ICS、CSV、Webhook、MQTT mock", "缓存、stale、重试和脱敏"],
            ["Renderer", "多 Profile、多语言、多状态", "像素 hash 与边界检查"],
            ["Compiler", "确定性、资源裁剪、签名和回滚", "同输入得到同 hash"],
            ["Device Contract", "检查更新、Range、激活确认、失败回退", "模拟器 + 真机矩阵"],
            ["Security", "SSRF、越权、重放、Secret 泄露", "发布阻断"],
            ["Performance", "编辑器、API、队列、编译和下载", "达到 P95 预算"],
        ],
        [1600, 4700, 3060],
    )

    d.heading("18. 实施阶段与启动计划", 1)
    d.table(
        ["阶段", "服务器端交付", "退出条件"],
        [
            ["S0 契约冻结", "DeviceProfile、Widget、Source、Display Package Schema", "平台和固件评审批准"],
            ["S1 基础平台", "账户、设备、Profile、项目、页面 CRUD、对象存储", "可创建 Profile 绑定项目"],
            ["S2 编辑与预览", "画布、Inspector、WASM 预览、预算", "400 x 300 BW/BWR 像素 fixture 通过"],
            ["S3 数据运行时", "REST/ICS/RSS/CSV、Snapshot、freshness、Vault", "标准数据源稳定运行"],
            ["S4 发布交付", "Compiler、签名、Package、设备 API、回滚", "NM-EPD-420 端到端激活"],
            ["S5 生态和运维", "Plugin registry、SDK、诊断、灰度和监控", "标准插件可独立升级"],
        ],
        [1500, 4800, 3060],
    )
    d.heading("18.1 启动团队建议", 2)
    d.bullets([
        "平台后端 2 人：领域 API、Connector、Compiler、设备交付。",
        "Web/编辑器 2 人：项目流程、画布、WASM 预览、数据绑定。",
        "固件 2 人：Runtime、包解析、渲染核心、驱动和低功耗。",
        "测试/DevOps 1 人：跨端 fixtures、设备农场、CI、环境和观测。",
        "产品/设计 1 人：模板、控件契约、数据源说明和评审推进。",
    ])

    d.heading("19. 主要风险与控制", 1)
    d.table(
        ["风险", "影响", "控制方案", "负责人"],
        [
            ["浏览器与设备渲染不一致", "页面重叠或截断", "共享 C++ 核心、字体 hash、像素 fixture", "Renderer"],
            ["Connector 数量增长", "限流、成本和故障扩散", "调度聚合、缓存、租户配额、隔离队列", "Data"],
            ["Schema 频繁变化", "包和固件不兼容", "版本矩阵、迁移器、ADR 和兼容 CI", "Architecture"],
            ["插件供应链", "恶意资源或权限滥用", "声明式包、签名、隔离构建、人工审批", "Security"],
            ["编辑器过度复杂", "延期和可用性下降", "先模板编辑，再开放自由画布", "Product"],
            ["签名密钥泄露", "全局制品信任失效", "KMS/HSM、最小权限、轮换和吊销", "Security"],
        ],
        [2200, 2350, 3400, 1410],
    )

    d.heading("20. 评审决策清单", 1)
    d.bullets([
        "是否批准模块化单体 + Worker 作为第一阶段服务形态。",
        "是否批准 TypeScript/Fastify/Next.js/PostgreSQL/Redis/S3 技术基线。",
        "是否批准可移植 C++17 渲染核心同时编译到 ESP32、native 和 WASM。",
        "是否批准设备配对采用 Device UUID + P-256，MAC 仅作诊断信息。",
        "是否批准平台只发布声明式插件包，第一阶段禁止设备执行插件代码。",
        "是否批准 S0-S5 实施顺序和第一阶段团队配置。",
    ])
    d.add_callout("启动门槛", "S0 结束前不得并行发明多个页面 JSON、字体格式或包协议。所有团队必须使用同一 contracts 仓库和 fixture 集。", COLOR_RED)
    d.save(SERVER_DOC)


def build_device_document(asset_dir: Path) -> None:
    layers = asset_dir / "device_layers.png"
    state_machine = asset_dir / "device_state_machine.png"
    create_device_layers(layers)
    create_device_state_machine(state_machine)

    d = ArchitectureDocument(
        "设备端架构设计",
        "EInk Dashboard Runtime、显示驱动、声明式控件、数据缓存与低功耗执行体系",
        "RB-EPD-ARCH-DEVICE-001",
        "ESP32 固件 Runtime、BSP/驱动、页面包、渲染、数据、功耗、存储、OTA 和测试",
    )
    d.add_cover()
    d.add_toc()
    add_common_document_control(d)

    d.heading("2. 目标、范围与架构定位", 1)
    d.paragraph("设备端定位为运行在 Arduino-ESP32/ESP-IDF、FreeRTOS 和板级 BSP 之上的通用墨水屏应用运行时。它不是新的操作系统内核，而是一层面向墨水屏的 Package、Application、Data、Graphics 和 System Runtime。")
    d.heading("2.1 目标", 2)
    d.bullets([
        "设备根据已签名 Display Package 渲染用户在服务器设计的多页 Dashboard。",
        "固件 Target 内置只读 Device Profile，并核验面板、驱动、字体和包兼容性。",
        "控件不直接访问 WiFi、NVS、LittleFS、GxEPD2 或其他控件内部状态。",
        "通过 active/candidate/previous 原子包机制、last-good Snapshot 和恢复界面实现离线可用。",
        "支持深度睡眠、短交互窗口、Focus Session、按键翻页、自动轮转和按页依赖同步。",
        "第一阶段支持 NM-EPD-420 400 x 300 黑白及黑白红两套独立 Profile。",
    ])
    d.heading("2.2 非目标", 2)
    d.bullets([
        "设备不承担页面设计、排序、主页编辑和插件市场管理。",
        "设备不执行服务器下发的任意脚本或动态原生代码。",
        "控件不自行猜测屏幕分辨率、颜色、局刷能力或可用内存。",
        "第一阶段不追求 LVGL 式连续动画、复杂触控和高帧率。",
        "不以单个全局字号或对整页位图缩放实现多设备适配。",
    ])

    d.heading("3. 分层架构与依赖规则", 1)
    d.figure(layers, "图 1 EInk Dashboard Runtime 分层")
    d.heading("3.1 分层职责", 2)
    d.table(
        ["层", "职责", "允许依赖"],
        [
            ["Package Runtime", "包校验、存储、激活、页面管理、调度", "Application/Data/System/Profile"],
            ["Application Runtime", "Widget、Binding、交互状态和页面应用逻辑", "Data/Graphics/System 抽象"],
            ["Data Runtime", "Snapshot、缓存、freshness、本地传感器", "System Services"],
            ["Graphics Runtime", "测量、布局、字体、颜色、合成、刷新计划", "Profile + Display 接口"],
            ["System Services", "网络、时间、输入、功耗、存储、OTA、诊断", "BSP/ESP-IDF/Arduino"],
            ["Hardware Abstraction", "板卡、面板、传感器和按键适配", "底层驱动"],
        ],
        [1900, 4300, 3160],
    )
    d.heading("3.2 强制依赖规则", 2)
    d.bullets([
        "Widget 只能读取 RuntimeContext、Snapshot 和资源句柄，只能向 DrawSurface 输出绘制命令。",
        "Layout/Font/Color 核心不包含 Arduino String、WiFiClient、GxEPD2 或 NVS 类型。",
        "Display Driver Adapter 是唯一可调用 GxEPD2 的模块。",
        "System Services 通过接口注入，不允许 Application Runtime 直接使用全局单例。",
        "跨层状态变更通过 RuntimeEvent 和显式命令发生，禁止回调中直接重启整套初始化。",
    ])

    d.heading("4. 技术基线与固件 Target", 1)
    d.table(
        ["领域", "技术基线", "说明"],
        [
            ["框架", "Arduino-ESP32 3.x / ESP-IDF 5.x", "保持当前 PlatformIO 路线，逐步使用 IDF 服务"],
            ["语言", "C++17", "禁止异常和 RTTI 的使用可按 Target 评估"],
            ["调度", "FreeRTOS + 单一 Runtime 事件队列", "长网络任务与渲染解耦"],
            ["显示", "GxEPD2 Driver Adapter", "页面/控件不依赖 GxEPD2"],
            ["配置", "ArduinoJson 7 + 有界解析器", "先校验大小、深度和版本"],
            ["文件系统", "LittleFS + 原子 manifest", "页面包、缓存、资源和诊断"],
            ["持久状态", "NVS + RTC memory", "设备设置、游标、计时和唤醒上下文"],
            ["安全", "mbedTLS + P-256 + CA bundle", "设备身份、HTTPS、签名和 hash"],
            ["构建", "PlatformIO 多环境", "每个 Target 绑定 BSP、Profile 和 Driver"],
        ],
        [1650, 3000, 4710],
    )
    d.heading("4.1 固件 Target", 2)
    d.code("""firmware_target: nm-display-420-bwr
board: nm-epd-420
profile_id: nm-epd-420-bwr-400x300-v1
driver_id: gxepd2.gdey042z98.bwr
driver_adapter_version: 1
runtime_schema_version: 1
widget_api_version: 1
font_pack_ids:
  - font.noto-inter.400x300.v1""")
    d.paragraph("黑白屏和黑白红屏必须是独立 Target/Profile/Driver Binding，即使板卡和分辨率相同，也不能只通过运行时颜色开关混用。")

    d.heading("5. Device Profile、BSP 与显示驱动", 1)
    d.heading("5.1 能力来源和核验", 2)
    d.code("""Platform Device Profile Registry
  -> freeze targetProfileId into project and package
  -> firmware target embeds read-only Profile Manifest
  -> BSP optionally probes panel ID / board revision / PSRAM
  -> Runtime reports match | mismatch | unavailable
  -> widgets never infer hardware capabilities""")
    d.heading("5.2 BSP 接口", 2)
    d.table(
        ["接口组", "建议能力", "实现要求"],
        [
            ["Board", "init, boardIdentity, hardwareProbe", "在显示初始化前完成板级电源和总线准备"],
            ["Display", "begin, createSurface, writePlane, refresh, sleep", "暴露能力，不暴露 GxEPD2 类型"],
            ["Input", "buttons, wake pins, debounce", "输出统一 ButtonEvent"],
            ["Power", "battery, rail control, deep/light sleep", "明确可唤醒引脚和关断顺序"],
            ["Sensor", "enumerate, sample, units", "时间戳、质量和错误状态标准化"],
            ["RTC", "read, set, validity, drift", "网络时间与 RTC 状态可解释"],
        ],
        [1750, 3300, 4310],
    )
    d.heading("5.3 IDisplayDriver", 2)
    d.code("""class IDisplayDriver {
 public:
  virtual DisplayCapabilities capabilities() const = 0;
  virtual Status begin(const DisplayInitContext&) = 0;
  virtual Surface createSurface(const SurfaceRequest&) = 0;
  virtual Status writePlane(ColorPlane, const BufferView&) = 0;
  virtual Status refreshFull() = 0;
  virtual Status refreshPartial(const RectList&) = 0;
  virtual Status sleep() = 0;
  virtual Status wake() = 0;
};""")
    d.heading("5.4 Driver Factory", 2)
    d.bullets([
        "每个 firmwareTarget 只注册编译并验证过的 Driver Adapter。",
        "页面包只引用 driverId，不能下载驱动代码。",
        "Adapter 版本改变刷新时序、颜色平面或像素格式时必须提升版本。",
        "真实面板测试覆盖全刷、局刷、红色平面、休眠唤醒和 busy 超时。",
    ])

    d.heading("6. RuntimeController 与运行模式", 1)
    d.figure(state_machine, "图 2 唤醒、渲染与休眠主状态机")
    d.heading("6.1 运行模式", 2)
    d.table(
        ["模式", "进入条件", "关键行为", "退出条件"],
        [
            ["Provisioning", "未配置网络或长按进入", "AP/配对/恢复；使用内置最小 UI", "保存配置并重启"],
            ["WakeRenderSleep", "定时或冷启动", "核验、按页同步、渲染、刷新、深睡", "提交状态后深睡"],
            ["InteractiveWindow", "按键唤醒或渲染后短窗口", "上一页/下一页、轮转设置", "空闲超时"],
            ["FocusSession", "focus_timer 长按启动", "WiFi 可关闭；分钟级更新；屏蔽翻页", "结束或长按停止"],
            ["Maintenance", "OTA、诊断或包恢复", "下载、校验、回滚、日志导出", "成功重启或安全回退"],
        ],
        [1600, 2200, 3500, 2060],
    )
    d.heading("6.2 事件模型", 2)
    d.code("""RuntimeEvent =
  WakeEvent | ButtonEvent | ScheduleEvent | PackageEvent |
  NetworkEvent | SnapshotEvent | FocusEvent | PowerEvent

RuntimeController::dispatch(event)
  -> validate current mode
  -> update bounded state
  -> enqueue command
  -> render only through RenderCommand
  -> persist only through StateCommitCommand""")
    d.heading("6.3 并发模型", 2)
    d.bullets([
        "RuntimeController 在单一控制任务中串行处理状态转换，避免页面、网络和按键同时修改全局状态。",
        "Network Worker 执行 HTTPS/MQTT，结果以 SnapshotEvent 返回；不得直接刷新屏幕。",
        "Render Worker 可独占显示和 scratch buffer；EPD busy 超时转为明确错误。",
        "Storage Worker 可用于大包写入，但 manifest 切换必须在控制任务中完成。",
        "所有队列长度和消息大小由 Device Profile/firmwareTarget 固定上限。",
    ])

    d.heading("7. Display Package 生命周期", 1)
    d.heading("7.1 存储槽位", 2)
    d.code("""active/      # currently trusted and rendered package
candidate/   # downloaded package under verification
previous/    # last known-good package
recovery/    # firmware built-in resources only""")
    d.heading("7.2 激活状态机", 2)
    d.code("""Downloaded
  -> HashVerified
  -> SignatureVerified
  -> ProfileVerified
  -> SchemaParsed
  -> ResourceChecked
  -> CandidateReady
  -> Activated
  -> FirstRenderVerified""")
    d.heading("7.3 验证顺序", 2)
    d.numbers([
        "检查下载长度、传输完整性和 SHA-256。",
        "验证签名 keyId、算法和吊销状态。",
        "匹配 profileId、profileSchemaVersion、firmwareTarget、driverId 和 Adapter 版本。",
        "匹配 runtimeSchemaVersion、widgetApiVersion、Font Pack 和 metrics hashes。",
        "使用流式/有界解析器检查页面、节点、嵌套、字符串和资源大小。",
        "检查所有 Widget、Source binding、字体、图标和页面引用。",
        "完成最小试解析和资源预算后，原子切换 manifest。",
    ])
    d.heading("7.4 回滚", 2)
    d.bullets([
        "candidate 任一步失败都不得覆盖 active。",
        "active 首次渲染失败时恢复 previous，并记录脱敏错误类别。",
        "active/previous 均不可用时显示固件内置恢复页面。",
        "恢复页面不依赖用户 Font Pack、网络 Snapshot 或第三方资源。",
    ])

    d.heading("8. Application Runtime 与 Widget 模型", 1)
    d.heading("8.1 页面对象", 2)
    d.code("""DisplayPageRef
  packageId
  packageVersion
  pageId        # stable business id
  pageNo        # generated display index

PageDocument
  metadata
  layoutVariant
  nodes[]
  bindings[]
  refreshPolicy
  empty/stale/error states""")
    d.heading("8.2 Widget 接口", 2)
    d.code("""class IWidgetRenderer {
 public:
  virtual WidgetMeasure measure(const MeasureContext&, const WidgetNode&) = 0;
  virtual Status layout(const LayoutContext&, WidgetNode&, const Rect&) = 0;
  virtual Status render(const RenderContext&, const WidgetNode&, DrawSurface&) = 0;
  virtual Status renderState(DataState, const WidgetNode&, DrawSurface&) = 0;
};""")
    d.heading("8.3 第一阶段 Widget", 2)
    d.table(
        ["类别", "基础 Widget", "复合/领域 Widget"],
        [
            ["通用", "text, icon, value, list, table, badge, divider", "chrome, empty_state"],
            ["时间", "digital_clock, date, countdown", "agenda_list, focus_timer, world_clock"],
            ["天气环境", "value, icon, sparkline, bar_chart", "weather_current, weekly_forecast, environment_summary"],
            ["任务内容", "checklist, list, progress", "habit_grid, headline_list, quote_card"],
            ["金融状态", "quote_value, sparkline, badge", "market_summary, service_status, build_status"],
        ],
        [1500, 3500, 4360],
    )
    d.heading("8.4 Focus Session 模块化", 2)
    d.paragraph("Focus 页面由可配置 focus_timer Widget 和固定 FocusSessionService 组成。用户可配置专注/休息时长、字体角色和布局，但不能修改计时、按键锁定和唤醒状态机。")
    d.bullets([
        "未激活时 USER 短按仍执行上一页，长按达到阈值立即启动，无需等待释放。",
        "激活后 USER 长按只唤醒 FocusSession 并停止，不进入网络初始化或切换 pageNo=0。",
        "倒计时按分钟显示，RTC/timer 保持秒级内部精度；仅跨分钟或状态变化时刷新。",
        "Focus 活动时屏蔽页面切换和自动轮转，结束后恢复正常 Runtime。",
    ])

    d.heading("9. Data Runtime、Snapshot 与按页同步", 1)
    d.heading("9.1 Snapshot 模型", 2)
    d.code("""SnapshotHeader
  sourceInstanceId
  schemaVersion
  snapshotVersion
  observedAt / fetchedAt / generatedAt / expiresAt
  sourceState: ready | stale | expired | error
  payloadLength / crc32

SnapshotPayload
  bounded typed values and arrays""")
    d.heading("9.2 数据获取策略", 2)
    d.bullets([
        "PageDependencyPlanner 从当前 PageDocument 计算 SourceInstance 和字段路径集合。",
        "启动、深睡唤醒和翻页只同步目标页面缺失或过期的依赖。",
        "平台执行源优先下载标准 Snapshot；device 执行源通过 ProviderRegistry 采集。",
        "网络失败时读取 last-good Snapshot，并依据 freshness 渲染 stale/expired/error。",
        "数据更新时间与 EPD 刷新时间分离；Snapshot 更新不等于立即全刷。",
    ])
    d.heading("9.3 Provider 边界", 2)
    d.table(
        ["接口", "职责", "禁止行为"],
        [
            ["fetch", "获取有限长度响应", "无超时、无限流读取"],
            ["parse", "结构化解析 JSON/CSV/RSS/ICS", "基于字符串搜索的脆弱解析"],
            ["normalize", "输出版本化 Snapshot", "暴露第三方原始字段给 Widget"],
            ["cache", "写入验证后的 last-good 数据", "失败响应覆盖缓存"],
            ["readiness", "解释 Required/Optional/Error", "以空白代替错误状态"],
        ],
        [1600, 3300, 4460],
    )

    d.heading("10. Graphics Runtime", 1)
    d.heading("10.1 渲染流水线", 2)
    d.code("""PageDocument + Snapshot + RuntimeContext
  -> Widget measure
  -> LayoutEngine native-pixel rectangles
  -> FontEngine glyph metrics and bitmap atlas
  -> DrawSurface logical commands
  -> ColorMapper physical planes
  -> Compositor and content hash
  -> RefreshPlanner
  -> IDisplayDriver""")
    d.heading("10.2 LayoutEngine", 2)
    d.bullets([
        "输入坐标是目标 Profile 的原生像素，支持绝对、锚点、网格、流式和受限断点布局。",
        "所有矩形使用有界整数；溢出、重叠和无效尺寸在发布时拦截，设备再次防御性校验。",
        "文本测量使用真实 glyph advance、bearing、baseline、line-height 和 fallback。",
        "不同分辨率使用不同 Font Profile 和布局变体，不通过全局缩放解决。",
    ])
    d.heading("10.3 FontEngine", 2)
    d.table(
        ["职责", "实现"],
        [
            ["语义角色", "title/body/caption/numericLarge/numericMedium/label 映射到确定字体资源"],
            ["免费字体", "Inter、Noto Sans CJK、Source Han Sans、Roboto Mono，版本和许可证随包"],
            ["字形裁剪", "发布器生成 glyph subset；设备只加载包所需块"],
            ["一致性", "fontPackSha256 和 fontMetricsSha256 激活前校验"],
            ["内存", "按页加载、LRU 或静态 atlas；不得将完整 CJK 字库常驻内部 Heap"],
        ],
        [2000, 7360],
    )
    d.heading("10.4 ColorMapper", 2)
    d.bullets([
        "Widget 只使用 Background、PrimaryText、MutedText、Accent、Warning、Critical 等逻辑颜色。",
        "BW Profile 把 Accent 映射为黑色或图案；BWR Profile 映射到红色平面。",
        "muted 不依赖透明度，通过网纹、轮廓、细线或 stale badge 表达。",
        "红色平面和黑色平面必须分离写入，禁止后写黑层覆盖红色符号。",
    ])
    d.heading("10.5 RefreshPlanner", 2)
    d.code("""RefreshDecision =
  Skip |
  PartialRefresh(rects) |
  FullRefresh |
  DeferUntil(nextWakeUtc)""")
    d.bullets([
        "综合 Profile 能力、颜色平面、差异矩形、ghosting、刷新计数、电池和运行模式。",
        "任何 Widget 不得直接请求全刷或局刷，只能声明内容变化和紧急程度。",
        "BWR 第一阶段可只支持全刷；BW Profile 可在真实面板验证后开放局刷。",
    ])

    d.heading("11. System Services", 1)
    d.table(
        ["服务", "主要接口", "关键规则"],
        [
            ["NetworkService", "connect, disconnect, status", "按需开启；退避；不保存默认生产 WiFi"],
            ["TimeService", "sync, nowUtc, timezone", "SNTP + RTC；标记未同步和漂移"],
            ["InputService", "events, hold, debounce", "按模式解释按键；长按达到阈值即触发"],
            ["PowerService", "planSleep, enterSleep, battery", "提交状态后休眠；验证唤醒引脚"],
            ["StorageService", "atomicWrite, slots, quota", "active/candidate/previous；掉电安全"],
            ["OtaService", "check, download, verify, switch", "签名、版本、回滚和 boot confirmation"],
            ["Diagnostics", "metrics, errors, export", "结构化、脱敏、有界环形日志"],
        ],
        [1750, 3100, 4510],
    )
    d.heading("11.1 时间与时区", 2)
    d.bullets([
        "内部统一使用 UTC epoch；渲染前依据设备 TimeZoneId 转本地时间。",
        "ICS floating + TZID 在展开前转换为 UTC，Windows 时区名通过紧凑映射表解析。",
        "跨周、跨月页面的日期范围每次渲染由 nowUtc 计算，不缓存固定 JUL 等文本。",
        "RTC 无效时显示未同步状态，不输出错误年份或时间。",
    ])

    d.heading("12. 存储与分区", 1)
    d.heading("12.1 建议分区职责", 2)
    d.table(
        ["区域", "内容", "一致性策略"],
        [
            ["NVS", "设备设置、网络、配对凭据、包游标、轮转、Focus 配置", "Schema 版本 + 双写迁移"],
            ["RTC Memory", "当前页面、唤醒上下文、Focus 运行状态、失败计数", "CRC/version；断电可丢失"],
            ["LittleFS package", "active/candidate/previous 页面包和资源", "临时文件 + fsync + manifest rename"],
            ["LittleFS cache", "Snapshot、HTTP metadata、last-good 数据", "按源配额和 LRU"],
            ["OTA", "双应用槽 + ota_data", "boot confirmation + rollback"],
            ["Core dump", "独立诊断分区", "Release 可禁用内容敏感 dump 或受控导出"],
        ],
        [1850, 4300, 3210],
    )
    d.heading("12.2 原子写入", 2)
    d.code("""write payload.tmp
  -> flush and close
  -> verify size/hash/schema
  -> write manifest.tmp
  -> rename payload.tmp to versioned payload
  -> rename manifest.tmp to candidate manifest
  -> activate by atomic active-manifest switch""")
    d.heading("12.3 配额", 2)
    d.bullets([
        "每个目录和 SourceInstance 有独立上限，避免单一缓存耗尽 LittleFS。",
        "包下载前检查剩余空间并预留 previous；空间不足先清理过期 cache，不删除 active。",
        "诊断日志采用固定大小环形缓冲，不允许无限增长。",
    ])

    d.heading("13. 内存与性能预算", 1)
    d.heading("13.1 内存原则", 2)
    d.bullets([
        "内部 Heap 保留给 FreeRTOS、网络栈、驱动和小对象；大 Snapshot、字体和 scratch buffer 优先 PSRAM。",
        "所有 JSON、数组、节点、字符串、图表点和下载长度有硬上限。",
        "渲染按页加载资源，不同时保留所有页面的场景树和 Snapshot。",
        "避免频繁 String 拼接和碎片化；使用 arena、span、string_view 和固定容量容器。",
    ])
    d.heading("13.2 NM-EPD-420 V1 预算基线", 2)
    d.table(
        ["资源", "建议预算", "门禁"],
        [
            ["用户页面", "16", "Profile 可调整，平台上限 64"],
            ["每页节点", "32", "解析前检查"],
            ["嵌套深度", "4", "Schema 硬限制"],
            ["图表点", "32/图表", "下采样由平台完成"],
            ["Display Package", "建议 256 KB", "下载前 Content-Length 检查"],
            ["Snapshot 单源", "按 SourceDefinition，建议 32-128 KB", "流式解析和配额"],
            ["内部 Heap 安全余量", "至少 64 KB", "渲染/网络阶段分别检查"],
            ["PSRAM", "8 MB 可用但不可无界", "峰值和最大连续块监控"],
        ],
        [2600, 3000, 3760],
    )

    d.heading("14. 按键、轮转与功耗", 1)
    d.heading("14.1 默认交互", 2)
    d.table(
        ["输入", "普通模式", "Focus 未激活", "Focus 激活"],
        [
            ["BOOT 短按", "下一页", "下一页", "屏蔽或仅唤醒"],
            ["USER 短按", "上一页", "上一页", "屏蔽"],
            ["USER 长按 2 秒", "页面定义动作或无动作", "立即启动 Focus", "立即停止 Focus"],
            ["定时唤醒", "按计划同步/渲染", "正常", "跨分钟更新倒计时"],
        ],
        [1700, 2500, 2500, 2660],
    )
    d.heading("14.2 深度睡眠", 2)
    d.numbers([
        "完成包/缓存和当前 pageId 游标提交。",
        "计算下一次页面、数据源、Focus 或维护唤醒时间。",
        "停止网络和外设任务，等待文件关闭。",
        "调用 BSP prepareForSleep，配置 BOOT/USER 可用的唤醒源。",
        "进入深度睡眠；唤醒后由 RuntimeController 根据 WakeEvent 恢复。",
    ])

    d.heading("15. 网络、安全与设备身份", 1)
    d.heading("15.1 设备身份", 2)
    d.bullets([
        "首次启动生成随机 Device UUID 和 P-256 密钥对；MAC 仅作诊断字段。",
        "私钥存储在受保护 NVS；量产版逐步启用 Secure Boot、Flash Encryption 和 NVS Encryption。",
        "配对使用一次性 challenge，平台签发短期、可轮换令牌。",
        "解绑和恢复出厂必须清除平台凭据、用户密钥和页面包。",
    ])
    d.heading("15.2 网络策略", 2)
    d.table(
        ["控制", "要求"],
        [
            ["TLS", "验证 CA、主机名和有效期；支持 CA bundle 更新"],
            ["HTTP", "超时、最大响应、Content-Length、chunked、gzip 策略明确"],
            ["URL", "不在普通日志输出 Secret URL；仅输出 host 和 hash"],
            ["重试", "指数退避 + jitter；区分 4xx、5xx、TLS、解析和限流"],
            ["OTA/Package", "hash + 签名 + 版本/Target/Profile 匹配"],
            ["本地 Portal", "仅设备配置与恢复；不提供页面设计写入接口"],
        ],
        [2200, 7160],
    )

    d.heading("16. 错误处理与恢复界面", 1)
    d.heading("16.1 错误分类", 2)
    d.table(
        ["类别", "例子", "设备行为"],
        [
            ["ProfileMismatch", "错误屏幕/Driver/Font Pack", "拒绝 candidate，保留 active"],
            ["PackageCorrupt", "hash、签名、解析失败", "删除 candidate 或隔离，记录错误码"],
            ["DataUnavailable", "网络/Provider 失败", "last-good + stale 或 empty_state"],
            ["RenderFailure", "资源不足、Widget 错误", "跳过问题节点或回滚 previous"],
            ["DisplayFailure", "busy 超时、总线错误", "有限重试，进入恢复页或休眠"],
            ["StorageFailure", "LittleFS/NVS 写入失败", "不切换 manifest，提示维护"],
        ],
        [2100, 3400, 3860],
    )
    d.heading("16.2 固件内置恢复页面", 2)
    d.bullets([
        "显示设备型号、固件版本、Profile、网络/配对状态和短错误码。",
        "提供重试页面包、进入配网、回滚 previous 和恢复出厂入口。",
        "只使用固件内置小字体和基础绘图，不依赖用户包。",
        "不显示 Secret、完整 URL、Token、日历内容或持仓数据。",
    ])

    d.heading("17. 诊断与可观测性", 1)
    d.table(
        ["指标", "采集位置", "用途"],
        [
            ["wake/render/sync duration", "RuntimeController", "定位启动和刷新延迟"],
            ["heap/PSRAM peak/max block", "Memory monitor", "发现碎片和预算偏差"],
            ["package validation result", "Package Runtime", "兼容和供应链诊断"],
            ["snapshot cache hit/stale", "Data Runtime", "网络与缓存质量"],
            ["full/partial/skip count", "RefreshPlanner", "面板寿命与功耗"],
            ["reset/wake cause", "System Service", "重启和按键问题"],
        ],
        [2700, 2700, 3960],
    )
    d.paragraph("日志采用模块、级别、错误码和关联 ID。生产日志不输出 WiFi 密码、API Key、Secret ICS URL、Token 或用户数据正文。诊断包需用户主动导出并经过脱敏。")

    d.heading("18. 测试策略", 1)
    d.table(
        ["测试层", "覆盖内容", "工具/门禁"],
        [
            ["Pure C++", "Schema、PageManager、Binding、freshness、布局、颜色、状态机", "PlatformIO native/host 测试"],
            ["Renderer fixture", "多语言、长文本、空/旧/错状态、BW/BWR", "像素 hash 和 bounds"],
            ["Driver contract", "颜色平面、busy、全刷/局刷、休眠", "fake driver + 真机"],
            ["Package fault", "截断、篡改、错 Profile、错 Font、掉电", "故障注入"],
            ["Network", "TLS、Range、ETag、超时、缓存、重试", "mock server"],
            ["Power/Input", "BOOT/USER、长按、深睡、Focus 分钟刷新", "GPIO fixture + COM 日志"],
            ["OTA", "升级、失败、回滚、旧包兼容", "双槽测试"],
            ["Soak", "24-72 小时轮转、唤醒和刷新", "真机矩阵与功耗记录"],
        ],
        [1600, 4700, 3060],
    )
    d.heading("18.1 真机验收矩阵", 2)
    d.bullets([
        "NM-EPD-420 BW：黑白色板、局刷能力按面板实测决定。",
        "NM-EPD-420 BWR：红色平面、全刷时间、休眠唤醒和长期 ghosting。",
        "400 x 300 多语言字体：中文、英文、数字、货币、单位和长标题。",
        "断网、弱网、平台离线、错误时间、低电量和文件系统接近满容量。",
    ])

    d.heading("19. 从当前代码迁移", 1)
    d.paragraph("当前仓库已经包含 PageManager、ProviderRegistry、RenderCoordinator、缓存、Calendar/Finance/News/Weather Provider、Web 配置和多个固定 renderer，可作为迁移素材，但 DashboardApp 已超过 1,900 行，需要按 Runtime 分层拆解。")
    d.table(
        ["当前模块", "目标模块", "迁移策略"],
        [
            ["DashboardApp", "RuntimeController + modes + commands", "先抽状态和事件，再移除直接编排"],
            ["IBoard/IEpdDriver", "BSP + IDisplayDriver + Factory", "扩展 capabilities 和 Driver Binding"],
            ["PageManager", "Package Runtime PageManager", "从固定 PageId 转为 package pageId/pageNo"],
            ["ProviderRegistry", "Data Runtime ProviderRegistry", "统一 Snapshot、readiness 和缓存契约"],
            ["RenderCoordinator", "RefreshPlanner", "加入 Profile、颜色平面、差异矩形和 ghosting"],
            ["固定 render_*", "Widget/Recipe", "逐页提取基础控件，完成后删除生产依赖"],
            ["WebServer 页面配置", "Local Device Portal", "保留 WiFi/配对/轮转/诊断，移除页面设计"],
            ["NVS AppConfig", "Device settings + package state", "版本化迁移，密钥继续独立存储"],
        ],
        [2300, 3100, 3960],
    )

    d.heading("20. 推荐代码目录", 1)
    d.code("""src/app/runtime/
  runtime_controller.*
  runtime_event.h
  runtime_mode.h
src/app/runtime/package/
  display_package.* package_store.* package_validator.*
  page_manager.* package_migration.*
src/app/runtime/application/
  ui_document.* widget_registry.* binding_resolver.*
  focus_session_service.*
src/app/runtime/data/
  data_snapshot.* snapshot_store.* freshness_policy.*
  local_sensor_bridge.*
src/app/runtime/graphics/
  layout_engine.* font_engine.* color_mapper.*
  compositor.* refresh_planner.*
src/app/profile/
  device_profile.* device_profile_registry.*
  driver_binding.* color_profile.* font_profile.*
src/bsp/display/
  display_driver.h display_driver_factory.* gxepd2_driver_adapter.*
src/ui/widgets/
  widget_text.* widget_clock.* widget_weather.* widget_calendar.* ...""")

    d.heading("21. 实施阶段与验收", 1)
    d.table(
        ["阶段", "设备端交付", "退出条件"],
        [
            ["D0 契约冻结", "Profile/Package/Widget/Snapshot 接口和 fixtures", "与平台 S0 同时批准"],
            ["D1 HAL 与状态机", "IDisplayDriver、Factory、RuntimeController、事件队列", "现有页面可通过适配层运行"],
            ["D2 包运行时", "active/candidate/previous、校验、恢复页", "故障注入不破坏 active"],
            ["D3 Graphics", "Layout/Font/Color/Refresh + native/WASM 共享核心", "BW/BWR 像素 fixture 通过"],
            ["D4 Data", "Snapshot、按页依赖、缓存、freshness", "断网仍显示可解释状态"],
            ["D5 Widget 迁移", "Clock/Weather/Agenda/Calendar/News/Finance/Focus", "固定 renderer 不再是生产路径"],
            ["D6 量产准备", "OTA、安全、功耗、soak、诊断和 Release", "真机矩阵和回滚通过"],
        ],
        [1500, 4800, 3060],
    )
    d.heading("21.1 设备端验收门禁", 2)
    d.bullets([
        "Profile/Driver/Font 任一不匹配时拒绝 candidate 且 active 不变。",
        "Pixel Accurate fixture 与服务器预览色板量化后像素 hash 一致。",
        "深睡唤醒只同步当前页依赖并恢复 packageId/pageId。",
        "Focus Active 下长按 USER 停止会话，不重连网络、不跳主页、不重启。",
        "所有网络和解析失败都可回退到缓存或明确 empty/error 状态。",
        "连续 72 小时轮转和睡眠测试无重启循环、文件系统损坏和内存持续下降。",
    ])

    d.heading("22. 评审决策清单", 1)
    d.bullets([
        "是否批准 Runtime 六层结构和单一 RuntimeController 事件模型。",
        "是否批准可移植 C++17 Graphics Core，不让 Widget 直接调用 GxEPD2。",
        "是否批准 BW/BWR 使用独立 firmwareTarget、Device Profile 和 Driver Binding。",
        "是否批准 active/candidate/previous + 内置恢复页的包生命周期。",
        "是否批准按页依赖同步、Snapshot freshness 和刷新解耦。",
        "是否批准 D0-D6 迁移顺序以及逐步拆解 DashboardApp 的方案。",
    ])
    d.add_callout("启动门槛", "D0 未冻结前，不开始大规模迁移固定页面。首个实现垂直切片应选择 Clock 或简单 Weather 页面，完整走通 Profile、Package、Snapshot、Widget、Renderer、Refresh 和 Sleep。", COLOR_RED)
    d.save(DEVICE_DOC)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="esp32_dashboard_arch_") as temp:
        asset_dir = Path(temp)
        build_server_document(asset_dir)
        build_device_document(asset_dir)
    print(SERVER_DOC)
    print(DEVICE_DOC)


if __name__ == "__main__":
    main()
